#!/usr/bin/env python3
"""LuaLaTeX 版の原稿から Word 版（.docx）を作る（経営情報 2026 AI利用ガイド 復習教材）。

原稿は LaTeX 版と共通。sessions/NN.tex（ワークシート）と answers/NN.tex（記入例）で
使っているマクロ（\\session、\\work、\\answerbox、\\abox、wstab/wstabh の表など）
だけを解釈し、Word の表と段落に置き換える。新しいマクロを原稿で使うときは、
ここにも対応を足すこと。基礎演習Ⅱのワークシート用スクリプトを流用している。

  python3 tex2docx.py -o word/worksheets.docx                 ワークシート全単元版（表紙つき）
  python3 tex2docx.py -o word/worksheet-03.docx 03            復習3のワークシートだけ
  python3 tex2docx.py --answers -o word/answers.docx          記入例の全単元版（表紙つき）
  python3 tex2docx.py --answers -o word/answers-03.docx 03    復習3の記入例だけ
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent

# ------------------------------------------------------------
# 体裁（preamble.tex に合わせる）
# ------------------------------------------------------------
MINCHO = "游明朝"
GOTHIC = "游ゴシック"
INK, INK2, INK3 = "1D2330", "4A5261", "757D8A"
RULE, PAPER2 = "B9BEB4", "ECEEE8"
BAD, BADTINT = "B3261E", "F9DEDA"
PHASES = {
    "1": ("1E4A4A", "ECF5D3", "AI利用ガイド　復習教材"),
}
ANS, ANSTINT = "C62828", "FBE9E7"  # 記入例の赤字と、解説の背景
COLORS = {"ans": ANS, "inkthree": INK3, "inktwo": INK2}
TEXT_W = 170.0  # 本文幅（mm）
SIZES = {"normal": 10, "small": 9, "footnotesize": 8, "scriptsize": 7, "large": 12}
FOOTER_LEFT = "経営情報 2026　AI利用ガイド　復習教材"


class Ctx:
    phase, tint = PHASES["1"][0], PHASES["1"][1]
    answers = False  # 記入例版のとき True


# ------------------------------------------------------------
# LaTeX の簡易パーサ
# ------------------------------------------------------------
def strip_comments(src):
    out = []
    for line in src.split("\n"):
        m = re.search(r"(?<!\\)%", line)
        out.append(line[: m.start()] if m else line)
    return "\n".join(out)


def match_close(s, i, op, cl):
    """s[i] が op のとき、対応する cl の位置を返す。"""
    depth = 0
    while i < len(s):
        c = s[i]
        if c == "\\":
            i += 2
            continue
        if c == op:
            depth += 1
        elif c == cl:
            depth -= 1
            if depth == 0:
                return i
        i += 1
    raise ValueError("括弧が閉じていません: " + s[:60])


def read_args(s, i, nopt, nreq):
    """位置 i から [opt] を最大 nopt 個、{arg} を nreq 個読む。"""
    opts, args = [], []
    for _ in range(nopt):
        j = i
        while j < len(s) and s[j] in " \t":
            j += 1
        if j < len(s) and s[j] == "[":
            k = match_close(s, j, "[", "]")
            opts.append(s[j + 1 : k])
            i = k + 1
        else:
            opts.append(None)
    for _ in range(nreq):
        while s[i] in " \t\n":
            i += 1
        if s[i] != "{":
            raise ValueError("引数がありません: " + s[i : i + 40])
        k = match_close(s, i, "{", "}")
        args.append(s[i + 1 : k])
        i = k + 1
    return opts, args, i


BLOCK_CMDS = {
    "session": (0, 6), "goals": (0, 1), "work": (1, 2), "instr": (0, 1),
    "answerbox": (1, 1), "twoboxes": (0, 3), "gridbox": (1, 1),
    "writelines": (0, 1), "fillin": (0, 1), "aimemo": (0, 0), "personacard": (0, 1),
    # 記入例版（answers-preamble.tex）のマクロ
    "abox": (1, 2), "atwoboxes": (0, 5), "alines": (0, 2), "afillin": (0, 2), "aimemoA": (0, 4),
}
IGNORED = {"noindent", "hfill", "par"}
ENVS = {
    "promptbox": (1, 0), "caution": (0, 1), "checks": (0, 0), "nexttime": (0, 0),
    "wstab": (0, 1), "wstabh": (0, 1), "minipage": (1, 1),
    "achecks": (0, 0), "point": (0, 0),
}


def parse_blocks(s):
    """ブロック単位のノード列にする。
    ('cmd', name, opts, args) / ('env', name, opts, args, body) / ('para', text)"""
    nodes, i, buf = [], 0, []

    def flush():
        t = "".join(buf).strip()
        if t:
            nodes.append(("para", t))
        buf.clear()

    while i < len(s):
        if s.startswith("\\begin{", i):
            flush()
            k = s.index("}", i)
            name = s[i + 7 : k]
            nopt, nreq = ENVS[name]
            opts, args, j = read_args(s, k + 1, nopt, nreq)
            end = find_env_end(s, j, name)
            nodes.append(("env", name, opts, args, s[j:end]))
            i = end + len("\\end{%s}" % name)
            continue
        m = re.match(r"\\([a-zA-Z]+)", s[i:])
        if m and m.group(1) in BLOCK_CMDS:
            flush()
            nopt, nreq = BLOCK_CMDS[m.group(1)]
            opts, args, j = read_args(s, i + m.end(), nopt, nreq)
            nodes.append(("cmd", m.group(1), opts, args))
            i = j
            continue
        if m and m.group(1) in IGNORED:
            if m.group(1) == "par":
                flush()
            i += m.end()
            continue
        if s.startswith("\n\n", i):
            flush()
        buf.append(s[i])
        i += 1
    flush()
    return nodes


def find_env_end(s, i, name):
    depth, b, e = 1, "\\begin{%s}" % name, "\\end{%s}" % name
    while True:
        nb, ne = s.find(b, i), s.find(e, i)
        if nb != -1 and nb < ne:
            depth, i = depth + 1, nb + len(b)
        else:
            depth -= 1
            if depth == 0:
                return ne
            i = ne + len(e)


def split_top(s, sep):
    """波括弧の外にある sep で分割する。"""
    parts, depth, cur, i = [], 0, [], 0
    while i < len(s):
        if s[i] == "\\" and not s.startswith(sep, i):
            cur.append(s[i : i + 2])
            i += 2
            continue
        if s[i] == "{":
            depth += 1
        elif s[i] == "}":
            depth -= 1
        if depth == 0 and s.startswith(sep, i):
            parts.append("".join(cur))
            cur, i = [], i + len(sep)
            continue
        cur.append(s[i])
        i += 1
    parts.append("".join(cur))
    return parts


def items_of(body):
    return [t.strip() for t in re.split(r"\\item\b", body)[1:] if t.strip()]


# ------------------------------------------------------------
# インライン（文字列 → 書式つきの断片）
# ------------------------------------------------------------
def inline(s, style=None):
    """[(text, style)] と改行 ('\n', None) の列を返す。"""
    style = dict(style or {})
    out, i = [], 0
    while i < len(s):
        c = s[i]
        if c == "{":
            k = match_close(s, i, "{", "}")
            out += inline(s[i + 1 : k], style)
            i = k + 1
            continue
        if c == "\\":
            m = re.match(r"\\([a-zA-Z]+|.)", s[i:])
            name, i = m.group(1), i + m.end()
            if name in SIZES:
                style["size"] = SIZES[name]
            elif name == "sffamily":
                style["gothic"] = True
            elif name == "bfseries":
                style["bold"] = True
            elif name == "newline":
                out.append(("\n", None))
            elif name == "underline":
                _, (a,), i = read_args(s, i, 0, 1)
                out += inline(a, dict(style, underline=True))
            elif name == "hspace":
                _, (a,), i = read_args(s, i, 0, 1)
                mm = float(re.match(r"[\d.]+", a).group())
                out.append(("　" * max(1, round(mm / 3.5)), dict(style)))
            elif name == "textbar":
                out.append(("｜", dict(style)))
            elif name == "quad":
                out.append(("　", dict(style)))
            elif name == "enspace":
                out.append((" ", dict(style)))
            elif name == "%":
                out.append(("%", dict(style)))
            elif name in (",", " "):
                pass
            elif name in ("hc",):
                _, (a,), i = read_args(s, i, 0, 1)
                out += inline(a, style)
            elif name == "ans":
                _, (a,), i = read_args(s, i, 0, 1)
                out += inline(a, dict(style, color=ANS, gothic=True))
            elif name == "af":
                _, (a,), i = read_args(s, i, 0, 1)
                out += inline(a, dict(style, color=ANS, gothic=True, size=8))
            elif name == "textbf":
                _, (a,), i = read_args(s, i, 0, 1)
                out += inline(a, dict(style, bold=True))
            elif name == "textcolor":
                _, (c, a), i = read_args(s, i, 0, 2)
                out += inline(a, dict(style, color=COLORS.get(c, INK)))
            elif name == "color":
                _, (c,), i = read_args(s, i, 0, 1)
                style["color"] = COLORS.get(c, INK)
            elif name == "makebox":
                _, (a,), i = read_args(s, i, 1, 1)
                out += [("\u3000", dict(style))] + inline(a, style) + [("\u3000", dict(style))]
            elif name == "achk":
                out.append(("☑", dict(style, color=ANS)))
            elif name == "par":
                out.append(("\n", None))
            else:
                raise ValueError("未対応のコマンド: \\" + name)
            continue
        j = i
        while j < len(s) and s[j] not in "{\\":
            j += 1
        text = re.sub(r"\s*\n\s*", "", s[i:j]).replace("~", "\u00a0")
        if text:
            out.append((text, dict(style)))
        i = j
    return out


def add_runs(p, frags, base=None):
    base = base or {}
    for text, st in frags:
        if text == "\n" and st is None:
            p.add_run().add_break(WD_BREAK.LINE)
            continue
        st = {**base, **st}
        r = p.add_run(text)
        font_of(r, st)


def font_of(r, st):
    name = GOTHIC if st.get("gothic") else MINCHO
    r.font.name = name
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), name)
    r.font.size = Pt(st.get("size", 10))
    r.font.bold = bool(st.get("bold"))
    if st.get("underline"):
        r.font.underline = True
    r.font.color.rgb = RGBColor.from_string(st.get("color", INK))


# ------------------------------------------------------------
# Word の部品
# ------------------------------------------------------------
def para(container, frags=(), base=None, align=None, before=0, after=0, keep=False, line=None):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line:
        pf.line_spacing = line
    if keep:
        pf.keep_with_next = True
    if align:
        p.alignment = align
    add_runs(p, frags, base)
    return p


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def border_xml(tag, spec):
    el = OxmlElement(tag)
    if spec is None:
        el.set(qn("w:val"), "nil")
        return el
    color, size = spec
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    return el


def set_cell_borders(cell, **sides):
    """sides: top/left/bottom/right = (color, 1/8pt) または None（線なし）。"""
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            b.append(border_xml("w:" + side, sides[side]))
    tcpr.append(b)


def set_cell_margins(cell, top=1.0, bottom=1.0, left=1.8, right=1.8):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, mm in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(int(mm * 56.7)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcpr.append(mar)


def new_table(container, ncols, widths_mm, nrows=1, borders=(RULE, 4)):
    t = container.add_table(rows=nrows, cols=ncols)
    t.autofit = False
    tbl = t._tbl
    tblpr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(widths_mm) * 56.7)))
    tw.set(qn("w:type"), "dxa")
    for old in tblpr.findall(qn("w:tblW")):
        tblpr.remove(old)
    tblpr.append(tw)
    tb = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tb.append(border_xml("w:" + side, borders))
    tblpr.append(tb)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_mm):
        gc.set(qn("w:w"), str(int(w * 56.7)))
    for row in t.rows:
        for cell, w in zip(row.cells, widths_mm):
            cell.width = Mm(w)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return t


def row_height(row, mm, exact=False):
    row.height = Mm(mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST
    cant = OxmlElement("w:cantSplit")
    row._tr.get_or_add_trPr().append(cant)


def cell_paragraph(cell, frags=(), base=None, align=None, keep=False):
    """セルの最初の段落（空）を使い、なければ足す。"""
    p = cell.paragraphs[0] if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text and not cell._tc.findall(qn("w:tbl")) and not getattr(cell, "_used", False) else cell.add_paragraph()
    cell._used = True
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    if keep:
        pf.keep_with_next = True
    if align:
        p.alignment = align
    add_runs(p, frags, base)
    return p


def keep_table_together(t):
    """最終行以外の段落に「次の段落と分離しない」を付け、表を1ページに収める。"""
    for row in t.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


def spacer(container, pt=3, keep=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.keep_with_next = keep
    pf.line_spacing = Pt(pt)
    r = p.add_run("")
    r.font.size = Pt(1)
    return p


def tidy_cell_end(cell):
    """入れ子の表の後に python-docx が足す空段落を目立たなくする。"""
    p = cell.paragraphs[-1]
    if not p.text:
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(1)


def drop_first_empty(cell):
    first = cell._tc.find(qn("w:p"))
    if first is not None and not "".join(first.itertext()):
        cell._tc.remove(first)


# ------------------------------------------------------------
# マクロごとの組み方
# ------------------------------------------------------------
def box(container, width, fill=None, border=(RULE, 4), height=None, exact=False, left_bar=None):
    t = new_table(container, 1, [width], borders=border)
    cell = t.cell(0, 0)
    if fill:
        set_cell_shading(cell, fill)
    if left_bar:
        set_cell_borders(cell, left=(left_bar, 24))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, 1.5, 1.5, 3, 3)
    if height:
        row_height(t.rows[0], height, exact)
    else:
        row_height(t.rows[0], 0)
    return t, cell


def name_line(container, fields, before=0, after=0, fills=None):
    """「学籍番号＿＿＿　氏名＿＿＿」の欄。fields = [(見出し, 下線の右端 mm)]、fills = {見出し: 記入例}"""
    p = para(container, [], before=before, after=after)
    st = {"gothic": True, "size": 9, "color": INK2}
    for k, (label, right) in enumerate(fields):
        p.paragraph_format.tab_stops.add_tab_stop(Mm(right), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES)
        font_of(p.add_run(("　" if k else "") + label), st)
        if fills and label in fills:
            font_of(p.add_run("　" + fills[label]), dict(st, color=ANS))
        font_of(p.add_run("\t"), st)
    return p


def label_style(size=7):
    return {"gothic": True, "size": size, "color": INK3}


def render_session_header(container, width, args):
    num, ph, date, title, tag, group = args
    Ctx.phase, Ctx.tint, phname = PHASES[ph]
    t, cell = box(container, width, fill=Ctx.phase, border=(Ctx.phase, 4))
    set_cell_margins(cell, 2.5, 2.5, 4, 4)
    sub = "復習%s　目安 %s　｜　%s" % (num, date, phname) + ("　｜　%s" % tag if tag.strip() else "")
    cell_paragraph(cell, [(sub, {})], {"gothic": True, "size": 9, "color": "FFFFFF"})
    cell_paragraph(cell, inline(title), {"gothic": True, "bold": True, "size": 18, "color": "FFFFFF"})
    fields = [("学籍番号", 50), ("氏名", 110)] + ([("グループ", 160)] if group.strip() else [])
    name_line(container, fields, before=4, after=4, fills={"氏名": "記入例"} if Ctx.answers else None)


def render_goals(container, width, body):
    t, cell = box(container, width, fill=Ctx.tint, border=(Ctx.tint, 4))
    cell_paragraph(cell, [("この回のねらい", {})], {"gothic": True, "bold": True, "size": 9, "color": Ctx.phase})
    for it in items_of(body):
        p = cell_paragraph(cell, [("・", {})] + inline(it), {"size": 9})
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-3.5)
    spacer(container)


def render_work(container, label, title):
    para(container, [("■ " + label, {"color": Ctx.phase}), ("　", {})] + inline(title, {}),
         {"gothic": True, "bold": True, "size": 10}, before=8, after=3, keep=True)


def render_answerbox(container, width, label, height_mm, grid=False):
    if grid:
        return render_gridbox(container, width, label, height_mm)
    t, cell = box(container, width, height=height_mm, exact=True)
    cell_paragraph(cell, inline(label or ""), label_style())
    spacer(container)


def render_gridbox(container, width, label, height_mm):
    if label:
        para(container, inline(label), label_style(), keep=True)
    step = 5.0
    ncols, nrows = int(width // step), int(height_mm // step)
    t = new_table(container, ncols, [width / ncols] * ncols, nrows=nrows, borders=("DDDDDD", 2))
    for row in t.rows:
        row_height(row, height_mm / nrows, exact=True)
        for cell in row.cells:
            set_cell_margins(cell, 0, 0, 0, 0)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = Pt(1)
    keep_table_together(t)
    spacer(container)


def render_twoboxes(container, width, a, b, h):
    side_by_side(container, width, [
        lambda c, w: render_answerbox(c, w, a, mm(h)),
        lambda c, w: render_answerbox(c, w, b, mm(h)),
    ])


def side_by_side(container, width, renderers, ratios=None):
    gap = width * 0.02
    n = len(renderers)
    each = (width - gap * (n - 1)) / n
    widths = []
    for k in range(n):
        widths.append(each)
        if k < n - 1:
            widths.append(gap)
    t = new_table(container, len(widths), widths, borders=None)
    row_height(t.rows[0], 0)
    for k, render in enumerate(renderers):
        cell = t.cell(0, 2 * k)
        set_cell_margins(cell, 0, 0, 0, 0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        drop_first_empty(cell)
        render(cell, each)
        tidy_cell_end(cell)
        if not cell.paragraphs:
            cell.add_paragraph()
    for k in range(1, len(widths), 2):
        set_cell_margins(t.cell(0, k), 0, 0, 0, 0)
    spacer(container)


def render_writelines(container, width, n):
    t = new_table(container, 1, [width], nrows=n, borders=None)
    for row in t.rows:
        row_height(row, 7.5, exact=True)
        set_cell_borders(row.cells[0], bottom=(RULE, 4))
    keep_table_together(t)
    spacer(container)


def render_fillin(container, width, label):
    p = para(container, inline(label), {"gothic": True, "size": 9}, before=2, after=6, keep=True)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(width - 0.5), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES)
    r = p.add_run("\t")
    font_of(r, {"size": 9, "color": INK3})


def render_instr(container, text):
    para(container, inline(text), {"size": 9, "color": INK2}, after=3, keep=True)


def render_prompt(container, width, title, body):
    t, cell = box(container, width, fill=PAPER2, border=(PAPER2, 4), left_bar=Ctx.phase)
    frags = inline(body.strip())
    cell_paragraph(cell, [("プロンプト例　" + (title or ""), {"bold": True, "size": 7, "color": Ctx.phase})] + [(" ", {})] + frags,
                   {"gothic": True, "size": 8, "color": INK2}, keep=True)
    spacer(container, 4, keep=True)


def render_caution(container, width, title, body):
    t, cell = box(container, width, fill=BADTINT, border=(BADTINT, 4))
    cell_paragraph(cell, [(title, {})], {"gothic": True, "bold": True, "size": 8, "color": BAD})
    cell_paragraph(cell, inline(body.strip()), {"size": 8})
    spacer(container, 4)


def render_checks(container, body):
    for it in items_of(body):
        p = para(container, [("□ ", {})] + inline(it), {"size": 9}, after=2)
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-4)


def render_nexttime(container, width, body):
    t, cell = box(container, width, border=(Ctx.phase, 6))
    cell_paragraph(cell, [("次回までに", {})], {"gothic": True, "bold": True, "size": 9, "color": Ctx.phase}, keep=True)
    for it in items_of(body):
        p = cell_paragraph(cell, [("□ ", {})] + inline(it), {"size": 9})
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-4)
    para(container, [], before=0, after=0).paragraph_format.keep_with_next = False
    spacer(container)


AIMEMO = r"""\begin{wstab}{colspec={Q[l,wd=30mm,m] X[l,m]},row{1-Z}={ht=9mm}}
    使用した場面 & \\
    主なプロンプト & \\
    出力をどう扱ったか\newline{\scriptsize 採用・修正・不採用と理由} & \\
    検証したこと\newline{\scriptsize 何で確かめたか} & \\
  \end{wstab}"""

PERSONA = r"""\begin{wstab}{colspec={Q[l,wd=26mm,m] X[l,m]},rows={ht=7.5mm}}
  \hc{ペルソナ#1　名前：} & \\
  年齢・職業 & \\ 家族構成 & \\ 居住地の特徴 & \\ 来店の目的 & \\ 重視する点 & \\
  不安や面倒 & \\ 休日の過ごし方 & \\
  \SetRow{ht=14mm}本人の言葉\newline\scriptsize 店舗に求めること & \\
\end{wstab}"""


def render_aimemo(container, width):
    para(container, [("■ AI活用記録（この単元の分）", {"color": Ctx.phase})], {"gothic": True, "bold": True}, before=8, after=3, keep=True)
    render_blocks(parse_blocks(AIMEMO), container, width)


# ------------------------------------------------------------
# 記入例版のマクロ（answers-preamble.tex）
# ------------------------------------------------------------
def split_pars(text):
    """\\par で段落に分ける。"""
    return [t.strip() for t in re.split(r"\\par\b", text) if t.strip()]


def render_abox(container, width, label, height_mm, content):
    t, cell = box(container, width, height=height_mm, exact=False)
    if label:
        cell_paragraph(cell, inline(label), label_style())
    for part in split_pars(content):
        cell_paragraph(cell, inline(part), {"gothic": True, "size": 9, "color": ANS})
    spacer(container)


def render_atwoboxes(container, width, a, b, h, ca, cb):
    side_by_side(container, width, [
        lambda c, w: render_abox(c, w, a, mm(h), ca),
        lambda c, w: render_abox(c, w, b, mm(h), cb),
    ])


def render_alines(container, width, n, content):
    """罫線に書き込んだ記入例。段落ごとに1行（長い段落は行の中で折り返す）。"""
    parts = split_pars(content)
    rows = max(n, len(parts))
    t = new_table(container, 1, [width], nrows=rows, borders=None)
    for k, row in enumerate(t.rows):
        row_height(row, 7.5)
        cell = row.cells[0]
        set_cell_borders(cell, bottom=(RULE, 4))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        if k < len(parts):
            cell_paragraph(cell, inline(parts[k]), {"gothic": True, "size": 9.5, "color": ANS})
    keep_table_together(t)
    spacer(container)


def render_afillin(container, width, label, content):
    p = para(container, inline(label) + [("　", {})] + inline(content, {"color": ANS}),
             {"gothic": True, "size": 9}, before=2, after=6, keep=True)
    ppr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    pb.append(border_xml("w:bottom", (INK3, 4)))
    ppr.append(pb)


def render_aimemoA(container, width, *fields):
    src = AIMEMO
    for f in fields:
        src = src.replace("& \\\\", "& \\af{%s} \\\\" % f, 1)
    para(container, [("■ AI活用記録（この単元の分）", {"color": Ctx.phase})], {"gothic": True, "bold": True}, before=8, after=3, keep=True)
    render_blocks(parse_blocks(src), container, width)


def render_achecks(container, body):
    for it in items_of(body):
        p = para(container, [("☑ ", {"color": ANS})] + inline(it), {"size": 9}, after=2)
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-4)


def render_point(container, width, body):
    t, cell = box(container, width, fill=ANSTINT, border=(ANSTINT, 4), left_bar=ANS)
    cell_paragraph(cell, [("解説・ポイント", {})], {"gothic": True, "bold": True, "size": 9, "color": ANS})
    for it in items_of(body):
        p = cell_paragraph(cell, [("・", {"color": ANS})] + inline(it), {"size": 9})
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-3.5)
    spacer(container)


def mm(s):
    return float(re.match(r"\s*([\d.]+)\s*mm", s).group(1))


def parse_colspec(spec):
    cols = []
    for kind, opts in re.findall(r"([QX])\[([^\]]*)\]", spec):
        o = [x.strip() for x in opts.split(",")]
        col = {"align": "c" if "c" in o else "l", "fixed": None, "weight": 1.0}
        for x in o:
            if x.startswith("wd="):
                col["fixed"] = mm(x[3:])
            elif re.fullmatch(r"[\d.]+", x):
                col["weight"] = float(x)
        cols.append(col)
    return cols


def render_table(container, width, spec, body, header):
    cs = re.search(r"colspec=\{(.*?)\}\s*(,|$)", spec, re.S)
    cols = parse_colspec(cs.group(1))
    fixed = sum(c["fixed"] or 0 for c in cols)
    wsum = sum(c["weight"] for c in cols if c["fixed"] is None) or 1
    widths = [c["fixed"] or (width - fixed) * c["weight"] / wsum for c in cols]

    heights = {}
    for sel, h in re.findall(r"row\{([\d\-Z]+)\}=\{[^}]*?ht=([\d.]+)mm", spec):
        heights[sel] = float(h)
    m = re.search(r"rows=\{ht=([\d.]+)mm", spec)
    all_h = float(m.group(1)) if m else None
    col1_bold = "column{1}" in spec

    rows = [r for r in split_top(body, r"\\") if r.strip()]
    rows = [split_top(r.strip(), "&") for r in rows]
    t = new_table(container, len(cols), widths, nrows=len(rows))
    for ri, cells in enumerate(rows):
        n = ri + 1
        h = heights.get(str(n)) or (heights.get("1-Z")) or (heights.get("2-Z") if n >= 2 else None) or all_h
        is_head = header and ri == 0
        row = t.rows[ri]
        for ci, raw in enumerate(cells[: len(cols)]):
            raw = raw.strip()
            m = re.match(r"\\SetRow\{ht=([\d.]+)mm\}", raw)
            if m:
                h, raw = float(m.group(1)), raw[m.end():]
            cell = row.cells[ci]
            base = {"size": 9}
            hc = raw.startswith("\\hc{")
            if is_head or hc:
                set_cell_shading(cell, Ctx.tint)
                base = {"gothic": True, "bold": True, "size": 8}
            if col1_bold and ci == 0:
                base = {"gothic": True, "bold": True, "size": 8}
            align = WD_ALIGN_PARAGRAPH.CENTER if (cols[ci]["align"] == "c" or is_head) else None
            cell_paragraph(cell, inline(raw), base, align)
        row_height(row, h or 0)
    keep_table_together(t)
    spacer(container, 4)


def render_blocks(nodes, container, width):
    i = 0
    while i < len(nodes):
        n = nodes[i]
        kind = n[0]
        if kind == "para":
            para(container, inline(n[1]), {"size": 9}, after=3)
        elif kind == "cmd":
            _, name, opts, args = n
            if name == "session":
                render_session_header(container, width, args)
            elif name == "goals":
                render_goals(container, width, args[0])
            elif name == "work":
                render_work(container, args[0], args[1])
            elif name == "instr":
                render_instr(container, args[0])
            elif name == "answerbox":
                render_answerbox(container, width, opts[0], mm(args[0]))
            elif name == "gridbox":
                render_gridbox(container, width, opts[0], mm(args[0]))
            elif name == "twoboxes":
                render_twoboxes(container, width, *args)
            elif name == "writelines":
                render_writelines(container, width, int(args[0]))
            elif name == "fillin":
                render_fillin(container, width, args[0])
            elif name == "aimemo":
                render_aimemo(container, width)
            elif name == "personacard":
                render_blocks(parse_blocks(PERSONA.replace("#1", args[0])), container, width)
            elif name == "abox":
                render_abox(container, width, opts[0], mm(args[0]), args[1])
            elif name == "atwoboxes":
                render_atwoboxes(container, width, *args)
            elif name == "alines":
                render_alines(container, width, int(args[0]), args[1])
            elif name == "afillin":
                render_afillin(container, width, args[0], args[1])
            elif name == "aimemoA":
                render_aimemoA(container, width, *args)
        elif kind == "env":
            _, name, opts, args, body = n
            if name == "minipage":
                group = []
                while i < len(nodes) and nodes[i][0] == "env" and nodes[i][1] == "minipage":
                    b = parse_blocks(nodes[i][4])
                    group.append(lambda c, w, b=b: render_blocks(b, c, w))
                    i += 1
                side_by_side(container, width, group)
                continue
            if name == "promptbox":
                render_prompt(container, width, opts[0], body)
            elif name == "caution":
                render_caution(container, width, args[0], body)
            elif name == "checks":
                render_checks(container, body)
            elif name == "nexttime":
                render_nexttime(container, width, body)
            elif name in ("wstab", "wstabh"):
                render_table(container, width, args[0], body, name == "wstabh")
            elif name == "achecks":
                render_achecks(container, body)
            elif name == "point":
                render_point(container, width, body)
        i += 1


# ------------------------------------------------------------
# 文書全体
# ------------------------------------------------------------
def add_field(run, instr):
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", instr), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:fldChar":
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = attr
        run._r.append(el)


def set_footer(section, right_text):
    section.footer.is_linked_to_previous = False
    f = section.footer
    p = f.paragraphs[0]
    p.style = None  # Footer スタイルの既定タブ（中央・右）を使わない
    p.text = ""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    st = {"gothic": True, "size": 7, "color": INK3}
    font_of(p.add_run(FOOTER_LEFT), st)
    if Ctx.answers:
        font_of(p.add_run("　記入例"), dict(st, color=ANS))
    if right_text is not None:
        font_of(p.add_run("\t" + right_text + "　"), st)
        r = p.add_run()
        font_of(r, st)
        add_field(r, "PAGE")


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Mm(20)
    s.top_margin, s.bottom_margin = Mm(16), Mm(16)
    s.footer_distance = Mm(8)
    normal = doc.styles["Normal"]
    normal.font.name = MINCHO
    normal.font.size = Pt(10)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), MINCHO)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:eastAsia"), "ja-JP")
    rpr.append(lang)
    pf = normal.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    # 既定の空段落を消す
    body = doc.element.body
    for p in body.findall(qn("w:p")):
        body.remove(p)


def session_meta(path):
    m = re.search(r"\\session\{(\d+)\}\{(\d)\}\{([^}]*)\}\{([^}]*)\}", path.read_text(encoding="utf-8"))
    return m.group(1), m.group(2), m.group(3), m.group(4)


def render_cover(doc, sessions):
    """表紙・使い方・目次。囲み（tcolorbox）の見出しと項目は cover.tex／answers-cover.tex から読む。"""
    src = strip_comments((HERE / ("answers-cover.tex" if Ctx.answers else "cover.tex")).read_text(encoding="utf-8"))
    para(doc, [], before=40)
    para(doc, [("経営情報 2026　AI利用ガイド", {})], {"gothic": True, "size": 12, "color": PHASES["1"][0]}, after=6)
    para(doc, [("復習教材", {})], {"gothic": True, "bold": True, "size": 30}, after=4)
    sub = [("問いを立て、調べ、読み、磨く", {})] + ([("　記入例", {"color": ANS})] if Ctx.answers else [])
    p = para(doc, sub, {"gothic": True, "size": 20}, after=10)
    ppr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    pb.append(border_xml("w:bottom", (INK2, 6)))
    ppr.append(pb)
    if not Ctx.answers:
        name_line(doc, [("学籍番号", 60), ("氏名", 140)], before=10, after=24)
    else:
        spacer(doc, 16)
    for k, m in enumerate(re.finditer(r"\\begin\{tcolorbox\}.*?\\end\{tcolorbox\}", src, re.S)):
        block = m.group(0)
        title = re.search(r"\\sffamily\\bfseries\s*([^}]*)\}", block).group(1).strip()
        items = items_of(re.search(r"\\begin\{itemize\}\[[^\]]*\](.*?)\\end\{itemize\}", block, re.S).group(1))
        if k == 0:
            t, cell = box(doc, TEXT_W, fill=PAPER2, border=(PAPER2, 4))
        else:
            t, cell = box(doc, TEXT_W)
        set_cell_margins(cell, 3, 3, 4, 4)
        cell_paragraph(cell, [(title, {})], {"gothic": True, "bold": True, "size": 9})
        for it in items:
            it = it.replace("PDFに直接書き込んでください", "このWordファイルに直接入力してください")
            p = cell_paragraph(cell, [("・", {})] + inline(it), {"size": 9})
            p.paragraph_format.left_indent = Mm(4)
            p.paragraph_format.first_line_indent = Mm(-3.5)
        spacer(doc, 8)
    para(doc, [("目次", {})], {"gothic": True, "bold": True}, before=12, after=4)
    t = new_table(doc, 2, [TEXT_W * 0.7, TEXT_W * 0.3], nrows=len(sessions), borders=None)
    for row, (num, _, date, title) in zip(t.rows, sessions):
        cell_paragraph(row.cells[0], [("復習%s　%s" % (num, title), {})], {"size": 9})
        cell_paragraph(row.cells[1], [("目安 " + date, {})], {"size": 9}, WD_ALIGN_PARAGRAPH.RIGHT)
        row_height(row, 5.5)
    set_footer(doc.sections[0], None)


# OOXML のスキーマが定める子要素の順（後から足した要素を並べ直す）
ORDER = {
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc "
             "tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd "
           "tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc "
           "textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
}


def normalize(root):
    """子要素をスキーマの順に並べ、同じ要素が重なったときは後のものを残す。"""
    for tag, names in ORDER.items():
        rank = {n: k for k, n in enumerate(names.split())}
        for el in root.iter(qn("w:" + tag)):
            kids = list(el)
            last = {}
            for kid in kids:
                last[kid.tag] = kid
            keep = [k for k in kids if last[k.tag] is k]
            keep.sort(key=lambda k: rank.get(k.tag.split("}")[1], len(rank)))
            for kid in kids:
                el.remove(kid)
            for kid in keep:
                el.append(kid)


def build(nums, out, cover):
    doc = Document()
    setup(doc)
    paths = [HERE / ("answers" if Ctx.answers else "sessions") / ("%s.tex" % n) for n in nums]
    metas = [session_meta(p) for p in paths]
    first = True
    if cover:
        render_cover(doc, metas)
        first = False
    for path, (num, ph, date, title) in zip(paths, metas):
        if first:
            section = doc.sections[0]
        else:
            section = doc.add_section(WD_SECTION.NEW_PAGE)
        set_footer(section, "復習%s　%s" % (num, title))
        first = False
        src = strip_comments(path.read_text(encoding="utf-8"))
        render_blocks(parse_blocks(src), doc, TEXT_W)
    normalize(doc.element)
    for sec in doc.sections:
        normalize(sec.footer._element)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "経営情報 2026 AI利用ガイド 復習教材" + (" 記入例" if Ctx.answers else "")
    doc.core_properties.author = "河合 勝彦"
    doc.save(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("nums", nargs="*", help="単元番号（2桁）。省略すると全単元版（表紙つき）")
    ap.add_argument("-o", "--output", required=True, type=Path)
    ap.add_argument("--answers", action="store_true", help="記入例版（answers/NN.tex）を作る")
    a = ap.parse_args()
    Ctx.answers = a.answers
    if a.nums:
        build(a.nums, a.output, cover=False)
    else:
        nums = sorted(p.stem for p in (HERE / "sessions").glob("*.tex"))
        build(nums, a.output, cover=True)


if __name__ == "__main__":
    main()
